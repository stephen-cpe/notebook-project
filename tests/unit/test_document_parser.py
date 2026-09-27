"""Unit tests for src.services.document_parser (TDD step 7).

Covers PDF/DOCX/PPTX/TXT/MD extraction against real fixture files:
- Each parser returns non-empty text for a valid fixture.
- TXT/MD parsers return exact content.
- PDF parser extracts multi-page text.
- DOCX parser extracts paragraph text.
- PPTX parser extracts slide text.
- Unknown/unsupported type raises a clear error.
- Missing file raises a clear error.
- Empty-text detection (for OCR fallback threshold).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src.services.document_parser import (
    detect_content_type,
    extract_text,
    parse_docx,
    parse_pdf,
    parse_pptx,
    parse_text_file,
)
from src.services.exceptions import IngestionError

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"


# ---------------------------------------------------------------------------
# Per-type parsing
# ---------------------------------------------------------------------------


class TestParsePdf:
    def test_extracts_text(self) -> None:
        text = parse_pdf(str(FIXTURES / "sample.pdf"))
        assert isinstance(text, str)
        assert "machine learning" in text.lower()
        assert "neural networks" in text.lower()

    def test_multi_page_content(self) -> None:
        text = parse_pdf(str(FIXTURES / "sample.pdf"))
        assert "page one" in text.lower()
        assert "page two" in text.lower()

    def test_returns_pages_count(self) -> None:
        from src.services.document_parser import parse_pdf_with_pages

        text, pages = parse_pdf_with_pages(str(FIXTURES / "sample.pdf"))
        assert pages == 2


class TestParsePdfDocument:
    def test_matches_legacy_helpers(self) -> None:
        """Single-read primitive must equal the dedicated passes exactly."""
        from src.services.document_parser import (
            extract_units,
            parse_pdf,
            parse_pdf_document,
            parse_pdf_pages,
            parse_pdf_with_pages,
        )

        path = str(FIXTURES / "sample.pdf")
        text, units, count = parse_pdf_document(path)
        assert text == parse_pdf(path)
        assert (text, count) == parse_pdf_with_pages(path)
        assert [(u.page, u.text) for u in units] == parse_pdf_pages(path)
        legacy_units, legacy_count = extract_units(path, "pdf")
        assert units == legacy_units
        assert count == legacy_count


class TestParseDocx:
    def test_extracts_text(self) -> None:
        text = parse_docx(str(FIXTURES / "sample.docx"))
        assert isinstance(text, str)
        assert "artificial intelligence" in text.lower()
        assert "transformers" in text.lower()

    def test_includes_headings(self) -> None:
        text = parse_docx(str(FIXTURES / "sample.docx"))
        assert "DOCX Fixture Document" in text


class TestParsePptx:
    def test_extracts_text(self) -> None:
        text = parse_pptx(str(FIXTURES / "sample.pptx"))
        assert isinstance(text, str)
        assert "cloud computing" in text.lower()
        assert "kubernetes" in text.lower()

    def test_includes_slide_titles(self) -> None:
        text = parse_pptx(str(FIXTURES / "sample.pptx"))
        assert "Slide One" in text or "Slide Two" in text


class TestParseTextFile:
    def test_txt_exact(self) -> None:
        text = parse_text_file(str(FIXTURES / "sample.txt"))
        assert "databases" in text.lower()
        assert "PostgreSQL" in text

    def test_md_exact(self) -> None:
        text = parse_text_file(str(FIXTURES / "sample.md"))
        assert "web development" in text.lower()
        assert "Flask" in text

    def test_empty_file_returns_empty(self, tmp_path: Path) -> None:
        p = tmp_path / "empty.txt"
        p.write_text("", encoding="utf-8")
        assert parse_text_file(str(p)) == ""

    def test_whitespace_only_returns_empty(self, tmp_path: Path) -> None:
        p = tmp_path / "ws.txt"
        p.write_text("   \n\n  \t  ", encoding="utf-8")
        assert parse_text_file(str(p)).strip() == ""


# ---------------------------------------------------------------------------
# detect_content_type
# ---------------------------------------------------------------------------


class TestDetectContentType:
    @pytest.mark.parametrize(
        ("filename", "expected"),
        [
            ("doc.pdf", "pdf"),
            ("report.PDF", "pdf"),
            ("file.docx", "docx"),
            ("slides.pptx", "pptx"),
            ("notes.txt", "txt"),
            ("readme.md", "md"),
            ("README.MD", "md"),
        ],
    )
    def test_extension_detection(self, filename: str, expected: str) -> None:
        assert detect_content_type(filename) == expected

    def test_unknown_extension_raises(self) -> None:
        with pytest.raises(IngestionError):
            detect_content_type("file.xyz")

    def test_no_extension_raises(self) -> None:
        with pytest.raises(IngestionError):
            detect_content_type("noextension")


# ---------------------------------------------------------------------------
# extract_text (dispatch by type)
# ---------------------------------------------------------------------------


class TestExtractText:
    def test_pdf(self) -> None:
        text = extract_text(str(FIXTURES / "sample.pdf"), "pdf")
        assert "machine learning" in text.lower()

    def test_docx(self) -> None:
        text = extract_text(str(FIXTURES / "sample.docx"), "docx")
        assert "artificial intelligence" in text.lower()

    def test_pptx(self) -> None:
        text = extract_text(str(FIXTURES / "sample.pptx"), "pptx")
        assert "cloud computing" in text.lower()

    def test_txt(self) -> None:
        text = extract_text(str(FIXTURES / "sample.txt"), "txt")
        assert "databases" in text.lower()

    def test_md(self) -> None:
        text = extract_text(str(FIXTURES / "sample.md"), "md")
        assert "web development" in text.lower()

    def test_unsupported_type_raises(self) -> None:
        with pytest.raises(IngestionError):
            extract_text(str(FIXTURES / "sample.txt"), "xyz")

    def test_missing_file_raises(self) -> None:
        with pytest.raises(IngestionError):
            extract_text("nonexistent_file_12345.pdf", "pdf")

    def test_empty_pdf_returns_empty(self) -> None:
        """A blank PDF (no text layer) returns empty -> triggers OCR fallback."""
        text = extract_text(str(FIXTURES / "empty.pdf"), "pdf")
        assert text.strip() == ""


# ---------------------------------------------------------------------------
# Image extraction (DOCX/PPTX) for OCR fallback
# ---------------------------------------------------------------------------


class TestExtractDocxImages:
    def test_extracts_image_from_docx_with_image(self) -> None:
        from src.services.document_parser import extract_docx_images

        images = extract_docx_images(str(FIXTURES / "_ocr_with_image.docx"))
        assert len(images) == 1
        # Each image is a PIL Image.
        from PIL import Image

        assert isinstance(images[0], Image.Image)

    def test_no_images_in_text_only_docx(self) -> None:
        from src.services.document_parser import extract_docx_images

        images = extract_docx_images(str(FIXTURES / "sample.docx"))
        assert images == []

    def test_missing_file_returns_empty(self) -> None:
        from src.services.document_parser import extract_docx_images

        assert extract_docx_images("nonexistent_12345.docx") == []


class TestExtractPptxImages:
    def test_extracts_image_from_pptx_with_image(self) -> None:
        from src.services.document_parser import extract_pptx_images

        images = extract_pptx_images(str(FIXTURES / "_ocr_with_image.pptx"))
        assert len(images) == 1
        from PIL import Image

        assert isinstance(images[0], Image.Image)

    def test_no_images_in_text_only_pptx(self) -> None:
        from src.services.document_parser import extract_pptx_images

        images = extract_pptx_images(str(FIXTURES / "sample.pptx"))
        assert images == []

    def test_missing_file_returns_empty(self) -> None:
        from src.services.document_parser import extract_pptx_images

        assert extract_pptx_images("nonexistent_12345.pptx") == []


class TestArchiveResourceBounds:
    def test_member_count_cap(self, tmp_path: Path) -> None:
        """An archive with too many media members yields no images."""
        import zipfile

        from src.services.document_parser import _extract_zip_media

        p = tmp_path / "many.docx"
        with zipfile.ZipFile(p, "w") as z:
            for i in range(101):
                z.writestr(f"word/media/img{i}.png", b"x")
        assert _extract_zip_media(str(p), "word/media/") == []

    def test_oversized_member_skipped(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Members over the per-member cap are skipped, not loaded."""
        import zipfile

        import src.services.document_parser as parser_mod
        from src.services.document_parser import _extract_zip_media

        p = tmp_path / "big.docx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/media/img0.png", b"y" * 64)
        monkeypatch.setattr(parser_mod, "_MAX_ARCHIVE_MEMBER_BYTES", 10)
        assert _extract_zip_media(str(p), "word/media/") == []

    def test_normal_archive_unaffected(self) -> None:
        """Ordinary fixtures still extract (bounds don't break real files)."""
        from src.services.document_parser import extract_docx_images

        images = extract_docx_images(str(FIXTURES / "_ocr_with_image.docx"))
        assert len(images) == 1


# ---------------------------------------------------------------------------
# Location-aware units (true pages/slides, None for pageless content)
# ---------------------------------------------------------------------------


class TestExtractUnits:
    def test_pdf_units_carry_real_page_numbers(self) -> None:
        from src.services.document_parser import extract_units

        units, page_count = extract_units(str(FIXTURES / "sample.pdf"), "pdf")
        assert page_count == 2
        assert len(units) >= 1
        pages = {u.page for u in units}
        assert pages <= {1, 2}
        assert 1 in pages

    def test_pptx_units_carry_slide_numbers(self) -> None:
        from src.services.document_parser import extract_units

        units, slide_count = extract_units(str(FIXTURES / "sample.pptx"), "pptx")
        assert slide_count is not None and slide_count >= 1
        assert len(units) >= 1
        for u in units:
            assert u.page is not None and 1 <= u.page <= slide_count

    def test_txt_unit_has_no_page(self) -> None:
        from src.services.document_parser import extract_units

        units, page_count = extract_units(str(FIXTURES / "sample.txt"), "txt")
        assert page_count is None
        assert len(units) == 1
        assert units[0].page is None

    def test_docx_tables_extracted(self, tmp_path: Path) -> None:
        """Table cell text must not be silently dropped."""
        from docx import Document

        from src.services.document_parser import extract_units, parse_docx

        p = tmp_path / "tabled.docx"
        doc = Document()
        doc.add_paragraph("Intro paragraph.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Alpha"
        table.cell(0, 1).text = "Beta"
        table.cell(1, 0).text = "Gamma"
        table.cell(1, 1).text = "Delta"
        doc.save(p)

        assert "Alpha" in parse_docx(str(p))
        assert "Delta" in parse_docx(str(p))
        units, page_count = extract_units(str(p), "docx")
        assert page_count is None
        assert any("Gamma" in u.text for u in units)
